"""
Atify tez markdown bölümlerini Düzce Üniversitesi Word şablonuna aktarır.

Çıktı: backend/tez/tez_atify_v1.docx

Kullanım:
    pip install python-docx markdown
    python build_docx.py

Şablon: C:/Users/atala/Desktop/tez_Düzce.docx (orijinal şablon, dokunulmaz)
Çıktı:  backend/tez/tez_atify_v1.docx (yeni dosya)

Bu script şu işi yapar:
- Şablonun bir kopyasını alır
- Yeni bölümler (heading + paragraf + çizelge + figür) ekler
- Mendeley ile sonradan kaynak gösterimi yapılması için yer tutucu
  bırakır.

NOT: Bu betik Word'ün otomatik içindekiler/numaralandırma alanlarını
yeniden oluşturmaz. Word'de aç, Ctrl+A → F9 ile alanları güncelle.
"""

from __future__ import annotations

import os
import re
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent.resolve()
TEMPLATE_SRC = Path("C:/Users/atala/Desktop/tez_Düzce.docx")
OUT_DOCX = HERE / "tez_atify_v1.docx"
FIGURES_DIR = HERE / "figures"

# Şablon yoksa kullanılacak yedek başlangıç noktaları
TEMPLATE_FALLBACKS = [
    HERE / "tez_atify_v1.docx",  # önceki üretilmiş çıktı
]

SECTIONS = [
    ("01_giris.md", "1. GİRİŞ"),
    ("02_literatur.md", "2. LİTERATÜR TARAMASI"),
    ("03_materyal_yontem.md", "3. MATERYAL VE YÖNTEM"),
    ("04_sistem_mimarisi.md", "4. SİSTEM MİMARİSİ VE TASARIMI"),
    ("05_algoritma.md", "5. AKUSTİK PARMAK İZİ ALGORİTMASI"),
    ("06_bulgular.md", "6. BULGULAR VE TARTIŞMA"),
    ("07_sonuclar.md", "7. SONUÇLAR VE ÖNERİLER"),
    ("08_kaynaklar.md", "8. KAYNAKLAR"),
]

# md başlık seviyesi → Word stili adı (şablon stilleri)
HEADING_STYLES = {
    1: "Heading 1",
    2: "Heading 2",
    3: "Heading 3",
    4: "Heading 4",
}


def add_heading(doc: Document, text: str, level: int) -> None:
    style_name = HEADING_STYLES.get(level, "Heading 4")
    try:
        para = doc.add_paragraph(text, style=style_name)
    except KeyError:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(16 - 2 * (level - 1))


def add_body(doc: Document, text: str) -> None:
    if not text.strip():
        return
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(11)


def add_table_from_md(doc: Document, header_row: list[str], data_rows: list[list[str]]) -> None:
    if not header_row:
        return
    table = doc.add_table(rows=1 + len(data_rows), cols=len(header_row))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for idx, h in enumerate(header_row):
        hdr_cells[idx].text = h.strip()
        for para in hdr_cells[idx].paragraphs:
            for run in para.runs:
                run.bold = True

    for r_idx, row in enumerate(data_rows, start=1):
        cells = table.rows[r_idx].cells
        for c_idx, cell_text in enumerate(row):
            if c_idx < len(cells):
                cells[c_idx].text = cell_text.strip()

    doc.add_paragraph()


def add_code_block(doc: Document, code: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(code.rstrip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def parse_md_table(lines: list[str], start_idx: int) -> tuple[list[str], list[list[str]], int] | None:
    """Markdown tablosunu parse eder. Geri dönüş: (header, rows, end_idx) veya None."""
    if start_idx >= len(lines):
        return None
    header_line = lines[start_idx].strip()
    if not (header_line.startswith("|") and header_line.endswith("|")):
        return None
    if start_idx + 1 >= len(lines):
        return None
    sep_line = lines[start_idx + 1].strip()
    if not re.match(r"^\|(\s*:?-+:?\s*\|)+$", sep_line):
        return None

    header = [c.strip() for c in header_line.strip("|").split("|")]
    rows = []
    i = start_idx + 2
    while i < len(lines):
        row = lines[i].strip()
        if not (row.startswith("|") and row.endswith("|")):
            break
        rows.append([c.strip() for c in row.strip("|").split("|")])
        i += 1
    return header, rows, i - 1


def render_markdown(doc: Document, md_text: str) -> None:
    lines = md_text.split("\n")
    i = 0
    para_buf: list[str] = []
    in_code = False
    code_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            text = " ".join(line.strip() for line in para_buf if line.strip())
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            text = re.sub(r"\[@[^\]]+\]", "[KAYNAK]", text)
            if text:
                add_body(doc, text)
            para_buf.clear()

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf.clear()
                in_code = False
            else:
                flush_para()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        m_h = re.match(r"^(#+)\s+(.+)$", line)
        if m_h:
            flush_para()
            level = len(m_h.group(1))
            text = m_h.group(2).strip()
            add_heading(doc, text, level)
            i += 1
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            flush_para()
            parsed = parse_md_table(lines, i)
            if parsed:
                header, rows, end_i = parsed
                add_table_from_md(doc, header, rows)
                i = end_i + 1
                continue

        if line.strip().startswith(">"):
            flush_para()
            quote = line.strip().lstrip(">").strip()
            quote = re.sub(r"\*\*(.+?)\*\*", r"\1", quote)
            quote = re.sub(r"\*(.+?)\*", r"\1", quote)
            para = doc.add_paragraph()
            run = para.add_run(quote)
            run.italic = True
            run.font.size = Pt(10)
            i += 1
            continue

        if re.match(r"^\s*[-*]\s+", line):
            flush_para()
            text = re.sub(r"^\s*[-*]\s+", "", line).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            add_body(doc, "•  " + text)
            i += 1
            continue

        m_num = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m_num:
            flush_para()
            text = m_num.group(2).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            add_body(doc, f"{m_num.group(1)}.  {text}")
            i += 1
            continue

        if not line.strip():
            flush_para()
            i += 1
            continue

        para_buf.append(line)
        i += 1

    flush_para()


# Bölüm numarasına göre o bölüme ait figürler (chapter -> [(file, caption)])
FIGURES_BY_CHAPTER: dict[int, list[tuple[str, str]]] = {
    2: [("sekil_2_1_prisma.png", "Şekil 2.1. PRISMA 2020 akış şeması")],
    4: [
        ("sekil_4_1_bilesen.png", "Şekil 4.1. Sistem bileşen şeması"),
        ("sekil_4_2_er_diyagrami.png", "Şekil 4.2. Veri tabanı ER diyagramı"),
        ("sekil_4_3_tanima_akisi.png", "Şekil 4.3. Şarkı tanıma sequence diyagramı"),
    ],
    5: [
        ("sekil_5_1_stft.png", "Şekil 5.1. STFT spektrogram örneği"),
        ("sekil_5_2_peaks.png", "Şekil 5.2. Bant bölmeli tepe seçimi"),
        ("sekil_5_3_landmarks.png", "Şekil 5.3. Anchor-target landmark eşleştirmesi"),
        ("sekil_5_4_histogram.png", "Şekil 5.4. Synthetic vs mikrofon ofset histogramı"),
    ],
    6: [("sekil_6_1_latency.png", "Şekil 6.1. Tanıma gecikme bileşenleri")],
}


def add_figure(doc: Document, fname: str, caption: str) -> None:
    path = FIGURES_DIR / fname
    if not path.exists():
        doc.add_paragraph(f"[{caption} — figür dosyası bulunamadı: {fname}]")
        return
    doc.add_picture(str(path), width=Cm(15.5))
    cap_para = doc.add_paragraph(caption)
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap_para.runs:
        run.italic = True
        run.font.size = Pt(10)
    doc.add_paragraph()


def add_chapter_figures(doc: Document, chapter_no: int) -> None:
    figs = FIGURES_BY_CHAPTER.get(chapter_no, [])
    for fname, caption in figs:
        add_figure(doc, fname, caption)


SCREENSHOTS = [
    ("ekran_anasayfa.png", "Şekil B.1. Atify ana sayfası (atify.com.tr)"),
    ("ekran_tanima.png", "Şekil B.2. Şarkı tanıma kayıt ekranı"),
    ("ekran_sonuc.png", "Şekil B.3. Tanıma sonuç ekranı"),
    ("ekran_admin.png", "Şekil B.4. Admin paneli"),
]


def add_page_number_footer(doc: Document) -> None:
    """Tüm bölümler için footer'a {PAGE} alanı ekler — sayfa numarası."""
    for section in doc.sections:
        footer = section.footer
        # Mevcut paragrafı temizle ve sayfa numarası ekle
        footer.is_linked_to_previous = False
        if footer.paragraphs:
            para = footer.paragraphs[0]
        else:
            para = footer.add_paragraph()
        # Mevcut içerikleri temizle
        for run in list(para.runs):
            run._element.getparent().remove(run._element)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = para.add_run()
        # PAGE field code'unu XML olarak ekle
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._element.append(fldChar1)

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE \\* MERGEFORMAT"
        run._element.append(instr)

        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run._element.append(fldChar2)


def strip_template_body(doc: Document) -> None:
    """
    Şablonun gövdesindeki tüm placeholder içeriği (Manzara şekilleri,
    Türkiye solar haritası, Birimler tablosu, kimyasal kısaltmalar,
    örnek bölümler) silip yalnızca header/footer/styles/sections'ı
    korur. Sayfa düzeni ve numbering bozulmasın diye XML'i doğrudan
    manipüle ediyoruz.
    """
    body = doc.element.body
    sectPr_list = []
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            sectPr_list.append(child)
            continue
        body.remove(child)
    # sectPr'ları geri ekle (sayfa boyutu/yönü için gerekli)
    for sp in sectPr_list:
        body.append(sp)


def main() -> None:
    source = TEMPLATE_SRC
    if not source.exists():
        for fb in TEMPLATE_FALLBACKS:
            if fb.exists() and fb != OUT_DOCX:
                source = fb
                break
        else:
            # Hiçbir şablon yok — çıktı dosyası varsa onu yeniden inşa için kullan
            if OUT_DOCX.exists():
                # Geçici kopyayı kaynak olarak kullan
                tmp_src = OUT_DOCX.with_name("_template_tmp.docx")
                shutil.copy(OUT_DOCX, tmp_src)
                source = tmp_src
            else:
                raise SystemExit(
                    "Sablon bulunamadi ve mevcut bir tez_atify_v1.docx da yok. "
                    "Sablonu C:/Users/atala/Desktop/tez_Düzce.docx adresine geri koyun."
                )

    print(f"Kaynak doc: {source}")
    if str(source) != str(OUT_DOCX):
        shutil.copy(source, OUT_DOCX)
    if source.name == "_template_tmp.docx":
        try:
            source.unlink()
        except Exception:
            pass

    doc = Document(str(OUT_DOCX))

    print("Mevcut icerigi temizliyorum...")
    strip_template_body(doc)

    # Kapak sayfasi
    print("  + Kapak")
    add_heading(doc, "T.C.", 2)
    add_heading(doc, "DÜZCE ÜNİVERSİTESİ", 2)
    add_heading(doc, "MÜHENDİSLİK FAKÜLTESİ", 2)
    add_heading(doc, "BİLGİSAYAR MÜHENDİSLİĞİ BÖLÜMÜ", 2)
    doc.add_paragraph()
    add_heading(doc,
        "AKUSTİK PARMAK İZİ YÖNTEMİ İLE ŞARKI TANIMA "
        "ÖZELLİĞİNE SAHİP WEB TABANLI MÜZİK AKIŞ PLATFORMU: ATIFY", 1)
    doc.add_paragraph()
    p = doc.add_paragraph("BM498 MEZUNİYET TEZİ")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph("Hazırlayan")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Atalay AKSOY")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("221002056")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph("Danışman")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Dr. Öğr. Üyesi Hüseyin BODUR")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph("Mayıs 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("DÜZCE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Beyan
    print("  + Beyan")
    add_heading(doc, "BEYAN", 1)
    add_body(doc,
        "Bu tez çalışmasının kendi çalışmam olduğunu, tezin planlanmasından "
        "yazımına kadar bütün aşamalarda etik dışı davranışımın olmadığını, "
        "bu tezdeki bütün bilgileri akademik ve etik kurallar içinde elde "
        "ettiğimi, bu tez çalışmasıyla elde edilmeyen bütün bilgi ve "
        "yorumlara kaynak gösterdiğimi ve bu kaynakları da kaynaklar listesine "
        "aldığımı, yine bu tezin çalışılması ve yazımı sırasında patent ve "
        "telif haklarını ihlal edici bir davranışımın olmadığını beyan ederim."
    )
    doc.add_paragraph()
    add_body(doc, "Mayıs 2026")
    add_body(doc, "(İmza)")
    add_body(doc, "Atalay AKSOY")
    doc.add_page_break()

    # UYZ
    print("  + Uretken Yapay Zeka Beyani")
    add_heading(doc, "ÜRETKEN YAPAY ZEKÂ KULLANIM BEYANI", 1)
    add_body(doc,
        "Bu tez çalışmasını hazırlarken ChatGPT (OpenAI) ve Claude (Anthropic) "
        "üretken yapay zekâ programlarından destek aldığımı beyan ederim. "
        "Tezimin hazırlığı aşamasında bu üretken yapay zekâ programlarından "
        "(i) dil çevirisi, (ii) bilimsel makalelere erişim, (iii) kaynak "
        "kodun Türkçe açıklanması ve (iv) akademik ifade kontrolü alanlarında "
        "destek aldım. Üretken yapay zekâ programlarından aldığım bilgilerin "
        "doğruluğunu kontrol ettiğimi bildiririm."
    )
    add_body(doc,
        "Herhangi bir zamanda, çalışmam ile ilgili yaptığım bu beyana aykırı "
        "bir durumun saptanması durumunda, ortaya çıkacak tüm ahlaki ve "
        "hukuki sonuçları kabul ettiğimi bildiririm."
    )
    doc.add_paragraph()
    add_body(doc, "Mayıs 2026")
    add_body(doc, "(İmza)")
    add_body(doc, "Atalay AKSOY")
    doc.add_page_break()

    # Tesekkur
    print("  + Tesekkur")
    add_heading(doc, "TEŞEKKÜR", 1)
    add_body(doc,
        "Lisans öğrenimim boyunca ve bu tezin hazırlanması sürecinde gösterdiği "
        "destek, sabır ve değerli yönlendirmeler için danışmanım Dr. Öğr. Üyesi "
        "Hüseyin BODUR'a en içten teşekkürlerimi sunarım."
    )
    add_body(doc,
        "Bu çalışma boyunca yardımlarını ve desteklerini esirgemeyen sevgili "
        "aileme ve çalışma arkadaşlarıma sonsuz teşekkürlerimi sunarım."
    )
    add_body(doc,
        "Atify projesinin geliştirme sürecinde açık lisanslı şarkı kataloğunu "
        "sağlayan Jamendo Music topluluğuna ve kullanılan tüm açık kaynak "
        "yazılım projelerine — Spring Boot, NumPy, MySQL, Docker, Caddy, "
        "ffmpeg — katkı sağlayan geliştiricilere minnettarım."
    )
    doc.add_page_break()

    # Icindekiler placeholder
    add_heading(doc, "İÇİNDEKİLER", 1)
    add_body(doc, "[Word'de Başvurular > İçindekiler menüsünden otomatik oluşturulacak]")
    doc.add_page_break()

    # Sekil/cizelge listeleri
    print("  + Sekil/Cizelge listeleri")
    add_heading(doc, "ŞEKİL LİSTESİ", 1)
    figures_for_list = [
        ("Şekil 2.1", "PRISMA 2020 akış şeması"),
        ("Şekil 4.1", "Sistem bileşen şeması"),
        ("Şekil 4.2", "Veri tabanı ER diyagramı"),
        ("Şekil 4.3", "Şarkı tanıma sequence diyagramı"),
        ("Şekil 5.1", "STFT spektrogram örneği"),
        ("Şekil 5.2", "Bant bölmeli tepe seçimi"),
        ("Şekil 5.3", "Anchor-target landmark eşleştirmesi"),
        ("Şekil 5.4", "Synthetic vs mikrofon ofset histogramı"),
        ("Şekil 6.1", "Tanıma gecikme bileşenleri"),
        ("Şekil B.1", "Atify ana sayfası"),
        ("Şekil B.2", "Şarkı tanıma kayıt ekranı"),
        ("Şekil B.3", "Tanıma sonuç ekranı"),
        ("Şekil B.4", "Admin paneli"),
    ]
    add_table_from_md(doc, ["No", "Açıklama"], figures_for_list)
    doc.add_page_break()

    add_heading(doc, "ÇİZELGE LİSTESİ", 1)
    tables_for_list = [
        ("Çizelge 3.1", "Backend yazılım bağımlılıkları"),
        ("Çizelge 3.2", "Recognizer servisi yazılım bağımlılıkları"),
        ("Çizelge 4.1", "song tablosunun seçili sütunları"),
        ("Çizelge 4.2", "Atify backend'in birincil REST uç noktaları"),
        ("Çizelge 4.3", "Recognizer servisi uç noktaları"),
        ("Çizelge 5.1", "STFT parametreleri"),
        ("Çizelge 5.2", "Sürüm 1 ile sürüm 2 arasındaki farklar"),
        ("Çizelge 6.1", "Seçili şarkılar için sürüm 2 hash sayıları"),
        ("Çizelge 6.2", "Tanıma akışı bileşenleri"),
        ("Çizelge 6.3", "Ters indeks öncesi/sonrası tanıma süresi karşılaştırması"),
        ("Çizelge 6.4", "recognition_attempt tablosu şeması"),
        ("Çizelge 7.1", "Tez amaçları ve gerçekleştirme durumu"),
    ]
    add_table_from_md(doc, ["No", "Açıklama"], tables_for_list)
    doc.add_page_break()

    # Kisaltmalar
    print("  + Kisaltmalar")
    add_heading(doc, "KISALTMALAR", 1)
    abbreviations = [
        ("API", "Application Programming Interface"),
        ("CD", "Continuous Delivery"),
        ("CDN", "Content Delivery Network"),
        ("CORS", "Cross-Origin Resource Sharing"),
        ("CPU", "Central Processing Unit"),
        ("DB", "Database"),
        ("DSP", "Digital Signal Processing"),
        ("FAIR", "Findable, Accessible, Interoperable, Reusable"),
        ("FFT", "Fast Fourier Transform"),
        ("HTTP", "Hypertext Transfer Protocol"),
        ("HTTPS", "HTTP Secure"),
        ("JPA", "Java Persistence API"),
        ("JSON", "JavaScript Object Notation"),
        ("JVM", "Java Virtual Machine"),
        ("JWT", "JSON Web Token"),
        ("MFCC", "Mel-Frequency Cepstral Coefficients"),
        ("ORM", "Object-Relational Mapping"),
        ("REST", "Representational State Transfer"),
        ("SPA", "Single Page Application"),
        ("STFT", "Short-Time Fourier Transform"),
        ("TLS", "Transport Layer Security"),
        ("VPS", "Virtual Private Server"),
        ("WAV", "Waveform Audio File Format"),
        ("WebM", "Web Media (acik kaynak medya konteyneri)"),
    ]
    add_table_from_md(doc, ["Kisaltma", "Acilim"], [[k, v] for k, v in abbreviations])
    doc.add_page_break()

    # Ozet
    print("  + Ozet/Abstract")
    abst_path = HERE / "09_ozet_abstract.md"
    if abst_path.exists():
        render_markdown(doc, abst_path.read_text(encoding="utf-8"))
    doc.add_page_break()

    # Bolumler
    for fname, title in SECTIONS:
        path = HERE / fname
        if not path.exists():
            print(f"  - Eksik: {fname}")
            continue
        print(f"  + {title}")
        md_text = path.read_text(encoding="utf-8")
        render_markdown(doc, md_text)
        # Bölüm numarasına göre figürler
        m = re.match(r"^(\d+)", title)
        if m:
            add_chapter_figures(doc, int(m.group(1)))
        doc.add_page_break()

    # EK B: Ekran goruntuleri
    print("  + Ek B: Ekran goruntuleri")
    add_heading(doc, "EK B: EKRAN GÖRÜNTÜLERİ", 1)
    add_body(doc,
        "Bu ekte, Atify platformunun üretim ortamından (atify.com.tr) "
        "alınan ekran görüntüleri sunulmuştur."
    )
    for fname, caption in SCREENSHOTS:
        path = FIGURES_DIR / fname
        if not path.exists():
            doc.add_paragraph(f"[{caption} — dosya bulunamadı: {fname}]")
            continue
        doc.add_picture(str(path), width=Cm(15.5))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(10)
        doc.add_paragraph()
    doc.add_page_break()

    # OZGECMIS
    print("  + Ozgecmis")
    add_heading(doc, "ÖZGEÇMİŞ", 1)
    add_heading(doc, "KİŞİSEL BİLGİLER", 2)
    add_table_from_md(doc, ["Alan", "Bilgi"], [
        ["Adı Soyadı", "Atalay AKSOY"],
        ["Öğrenci No", "221002056"],
        ["E-posta", "atalay-aksoy@hotmail.com"],
        ["Yabancı Dili", "İngilizce"],
    ])
    add_heading(doc, "ÖĞRENİM DURUMU", 2)
    add_table_from_md(doc, ["Derece", "Alan", "Okul/Üniversite", "Mezuniyet Yılı"], [
        ["Lisans", "Bilgisayar Mühendisliği", "Düzce Üniversitesi", "2026"],
    ])
    add_heading(doc, "PROJELER", 2)
    add_body(doc,
        "Atify (2025-2026): Akustik parmak izi tabanlı şarkı tanıma "
        "özelliğine sahip web tabanlı müzik akış platformu. Spring Boot, "
        "Python, MySQL, Docker, Caddy. Üretim ortamı: https://atify.com.tr"
    )

    # Sayfa numaralari (footer)
    print("  + Sayfa numaralari")
    add_page_number_footer(doc)

    doc.save(str(OUT_DOCX))
    print(f"\nKayit edildi: {OUT_DOCX}")
    print("Word'de acip Ctrl+A -> F9 ile alanlari guncelle.")


if __name__ == "__main__":
    main()
